"""
Extracts raw text from uploaded resume files (PDF or DOCX).
"""
import pdfplumber
import docx


def extract_text_from_pdf(file_path):
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_path):
    document = docx.Document(file_path)
    text_parts = [para.text for para in document.paragraphs if para.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    return "\n".join(text_parts)


def extract_text(file_path):
    """Dispatch based on file extension."""
    ext = str(file_path).lower().split('.')[-1]
    if ext == 'pdf':
        return extract_text_from_pdf(file_path)
    elif ext == 'docx':
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")